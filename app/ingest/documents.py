"""文档解析：pdf / docx / xlsx / pptx / csv / txt → 纯文本（FR-ING-02）。"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from app.util import clamp


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".xlsx":
        return _xlsx(data)
    if ext == ".pptx":
        return _pptx(data)
    if ext in (".csv",):
        return _csv(data)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"不支持的文档类型: {ext}")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = [(page.extract_text() or "") for page in reader.pages[:50]]
    return clamp("\n".join(parts), 20000)


def _docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables[:20]:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return clamp("\n".join(parts), 20000)


def _xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets[:10]:
        lines.append(f"# sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 500:
                lines.append("...（截断）")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                lines.append("\t".join(vals))
    return clamp("\n".join(lines), 20000)


def _pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"# slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    lines.append(t)
    return clamp("\n".join(lines), 20000)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))[:500]
    return clamp("\n".join("\t".join(r) for r in rows), 20000)
